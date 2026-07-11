"""One-time import: Community Resources xlsx -> resources.db.

Normalizes the per-sheet header variants into one schema (see db.py).
The 'Homeless Education' sheet has its own layout and is mapped specially.

Usage:
    python import_xlsx.py "Community Resources 2025 Jul- verified (final, all)).xlsx"
"""
import re
import sys
from datetime import datetime

import openpyxl

from db import get_conn, init_db

# The workbook filename says "2025 Jul- verified", so every imported row
# starts life verified as of July 2025.
IMPORT_VERIFIED_DATE = "2025-07-01"
IMPORT_VERIFIED_NOTE = "Imported from Excel (verified July 2025)"

# Header -> canonical field. Headers are matched after lowercasing and
# collapsing whitespace, so the newline-wrapped headers match too.
HEADER_MAP = {
    "type": "type",
    "resource type": "type",
    "area": "area",
    "organization name": "org_name",
    "contact (phone)": "phone",
    "contact information": "phone",
    "phone number": "phone",
    "phone": "phone",
    "contact (email)": "email",
    "email": "email",
    "e-mail": "email",
    "location": "location",
    "hours": "hours",
    "website": "website",
    "notes": "notes",
    "other notes": "notes",
    "bilingual language support?": "bilingual_support",
    "accessibility": "accessibility",
    "insurance/cost": "insurance_cost",
    "specific eligibility criteria (visit the website for other eligibility criteria)": "eligibility_criteria",
    "income based eligibility? (visit website for more info)": "income_based",
    "flyer?": "has_flyer",
    "flyer": "has_flyer",
    "application form": "has_application_form",
}

BOOL_FIELDS = {"income_based", "has_flyer", "has_application_form"}

# Source-data cleanup for the area column: it feeds the UI filter dropdown,
# so spelling/case variants would show up as separate areas.
AREA_FIXES = {
    "Thronton": "Thornton",
    "Not specified": "Not Specified",
}


def norm_header(value):
    if value is None:
        return None
    return re.sub(r"\s+", " ", str(value)).strip().lower()


def clean(value):
    """Cell value -> stripped string or None."""
    if value is None:
        return None
    # _x000D_ is how Excel escapes stray carriage returns inside cells
    s = str(value).replace("_x000D_", "\n").replace("\r", "\n").strip()
    return s or None


def to_bool(value):
    if value is None:
        return None
    if isinstance(value, bool):
        return int(value)
    s = str(value).strip().lower()
    if not s:
        return None
    if s in ("yes", "y", "true", "1") or s.startswith("yes"):
        return 1
    if s in ("no", "n", "false", "0") or s.startswith("no"):
        return 0
    return None  # free text we can't call either way


def rows_from_standard_sheet(ws):
    headers = [norm_header(c.value) for c in next(ws.iter_rows(min_row=1, max_row=1))]
    fields = [HEADER_MAP.get(h) for h in headers]
    unmapped = [h for h, f in zip(headers, fields) if h and not f]
    if unmapped:
        print(f"  ! unmapped headers in '{ws.title}': {unmapped}")

    for row in ws.iter_rows(min_row=2, values_only=True):
        record = {}
        for field, value in zip(fields, row):
            if field is None:
                continue
            record[field] = to_bool(value) if field in BOOL_FIELDS else clean(value)
        if record.get("area") in AREA_FIXES:
            record["area"] = AREA_FIXES[record["area"]]
        if record.get("org_name"):
            yield record


def rows_from_homeless_education(ws):
    """Columns: School District, District Legal Liaison, E-mail, Phone, Alternate Contact."""
    for row in ws.iter_rows(min_row=2, values_only=True):
        district, liaison, email, phone, alternate = (clean(v) for v in row[:5])
        if not district:
            continue
        # Area is usually in parentheses, e.g. "Adams 12 (Northglenn/Thornton)"
        m = re.search(r"\(([^)]+)\)", district)
        notes = []
        if liaison:
            notes.append(f"District Legal Liaison: {liaison}")
        if alternate:
            notes.append(f"Alternate contact: {alternate}")
        yield {
            "org_name": district,
            "type": "Homeless Education Liaison (McKinney-Vento)",
            "area": m.group(1) if m else None,
            "email": email,
            "phone": phone,
            "notes": "\n".join(notes) or None,
        }


def run(xlsx_path):
    init_db()
    wb = openpyxl.load_workbook(xlsx_path, read_only=True)
    conn = get_conn()
    existing = conn.execute("SELECT COUNT(*) FROM resources").fetchone()[0]
    if existing:
        sys.exit(f"resources.db already has {existing} rows - not importing again. "
                 f"Delete resources.db first to re-import.")

    now = datetime.now().isoformat(timespec="seconds")
    total = 0
    for sheet in wb.sheetnames:
        ws = wb[sheet]
        rows = (rows_from_homeless_education(ws) if sheet == "Homeless Education"
                else rows_from_standard_sheet(ws))
        count = 0
        for record in rows:
            record.update(
                category=sheet,
                status="active",
                last_verified_date=IMPORT_VERIFIED_DATE,
                verification_note=IMPORT_VERIFIED_NOTE,
                created_at=now,
                updated_at=now,
            )
            cols = ", ".join(record)
            marks = ", ".join("?" for _ in record)
            conn.execute(f"INSERT INTO resources ({cols}) VALUES ({marks})",
                         list(record.values()))
            count += 1
        total += count
        print(f"  {sheet}: {count} resources")
    conn.commit()
    conn.close()
    print(f"Imported {total} resources.")


if __name__ == "__main__":
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    run(sys.argv[1])
