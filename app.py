"""Community Resource Navigator.

Phase 1: searchable/editable resource list + Excel import/export.
Phase 2: AI verification agent + review queue (see verify_agent.py).
"""
import io
import json
import secrets
import threading
from datetime import date, datetime

import openpyxl
from openpyxl.styles import Font
from flask import Flask, jsonify, render_template, request, send_file

import matcher
import verify_agent
from db import get_conn, init_db

app = Flask(__name__)
init_db()

SEARCH_FIELDS = ("org_name", "type", "area", "location", "notes",
                 "eligibility_criteria", "phone", "email", "website")

EDITABLE_FIELDS = ("category", "type", "area", "org_name", "phone", "email",
                   "location", "hours", "website", "notes", "bilingual_support",
                   "accessibility", "insurance_cost", "eligibility_criteria",
                   "income_based", "has_flyer", "has_application_form",
                   "status", "seasonal_note", "last_verified_date",
                   "verification_note")

EXPORT_COLUMNS = ("id", "category", "type", "area", "org_name", "phone",
                  "email", "location", "hours", "website", "notes",
                  "bilingual_support", "accessibility", "insurance_cost",
                  "eligibility_criteria", "income_based", "has_flyer",
                  "has_application_form", "status", "seasonal_note",
                  "last_verified_date", "verification_note")


@app.get("/")
def index():
    return render_template("index.html")


@app.get("/api/resources")
def list_resources():
    where, params = [], []
    q = request.args.get("q", "").strip()
    if q:
        like = f"%{q}%"
        where.append("(" + " OR ".join(f"{f} LIKE ?" for f in SEARCH_FIELDS) + ")")
        params += [like] * len(SEARCH_FIELDS)
    for field in ("category", "area", "status"):
        value = request.args.get(field, "").strip()
        if value:
            where.append(f"{field} = ?")
            params.append(value)
    sql = "SELECT * FROM resources"
    if where:
        sql += " WHERE " + " AND ".join(where)
    sql += " ORDER BY category, org_name COLLATE NOCASE"
    conn = get_conn()
    rows = [dict(r) for r in conn.execute(sql, params)]
    conn.close()
    return jsonify(rows)


@app.get("/api/meta")
def meta():
    conn = get_conn()
    categories = [r[0] for r in conn.execute(
        "SELECT DISTINCT category FROM resources ORDER BY category")]
    areas = [r[0] for r in conn.execute(
        "SELECT DISTINCT area FROM resources WHERE area IS NOT NULL ORDER BY area")]
    total, active = conn.execute(
        "SELECT COUNT(*), SUM(status = 'active') FROM resources").fetchone()
    verified_90d = conn.execute(
        "SELECT COUNT(*) FROM resources WHERE last_verified_date >= date('now', '-90 days')"
    ).fetchone()[0]
    pending_review = conn.execute(
        "SELECT COUNT(*) FROM verification_checks WHERE review_status = 'pending'"
    ).fetchone()[0]
    conn.close()
    return jsonify({"categories": categories, "areas": areas, "total": total,
                    "active": active or 0, "verified_90d": verified_90d,
                    "pending_review": pending_review})


def validated_payload(data):
    payload, errors = {}, []
    for field in EDITABLE_FIELDS:
        if field not in data:
            continue
        value = data[field]
        if isinstance(value, str):
            value = value.strip() or None
        if field in ("income_based", "has_flyer", "has_application_form") and value is not None:
            value = 1 if value in (1, True, "1", "true", "yes") else 0
        if field == "status" and value not in ("active", "needs_review", "inactive"):
            errors.append("status must be active, needs_review or inactive")
        payload[field] = value
    if "org_name" in payload and not payload["org_name"]:
        errors.append("Organization name is required")
    if "category" in payload and not payload["category"]:
        errors.append("Category is required")
    return payload, errors


@app.post("/api/resources")
def create_resource():
    payload, errors = validated_payload(request.get_json(force=True))
    if not payload.get("org_name"):
        errors.append("Organization name is required")
    if not payload.get("category"):
        errors.append("Category is required")
    if errors:
        return jsonify({"errors": sorted(set(errors))}), 400
    payload.setdefault("status", "active")
    now = datetime.now().isoformat(timespec="seconds")
    payload["created_at"] = payload["updated_at"] = now
    conn = get_conn()
    cols = ", ".join(payload)
    marks = ", ".join("?" for _ in payload)
    cur = conn.execute(f"INSERT INTO resources ({cols}) VALUES ({marks})",
                       list(payload.values()))
    conn.commit()
    row = dict(conn.execute("SELECT * FROM resources WHERE id = ?",
                            (cur.lastrowid,)).fetchone())
    conn.close()
    return jsonify(row), 201


@app.put("/api/resources/<int:rid>")
def update_resource(rid):
    payload, errors = validated_payload(request.get_json(force=True))
    if errors:
        return jsonify({"errors": sorted(set(errors))}), 400
    if not payload:
        return jsonify({"errors": ["No editable fields in request"]}), 400
    payload["updated_at"] = datetime.now().isoformat(timespec="seconds")
    conn = get_conn()
    sets = ", ".join(f"{f} = ?" for f in payload)
    cur = conn.execute(f"UPDATE resources SET {sets} WHERE id = ?",
                       list(payload.values()) + [rid])
    conn.commit()
    if cur.rowcount == 0:
        conn.close()
        return jsonify({"errors": ["Resource not found"]}), 404
    row = dict(conn.execute("SELECT * FROM resources WHERE id = ?", (rid,)).fetchone())
    conn.close()
    return jsonify(row)


# ---------- Phase 2: verification + review queue ----------

VERIFY_STATE = {"running": False, "done": 0, "total": 0, "current": None,
                "verdicts": {}, "started_at": None, "finished_at": None,
                "error": None, "stopped": None}
_verify_lock = threading.Lock()


def _run_verification(limit, category):
    try:
        for summary in verify_agent.run(limit=limit, category=category):
            VERIFY_STATE.update(done=summary["done"], total=summary["total"],
                                current=summary["current"],
                                verdicts=summary["verdicts"],
                                stopped=summary.get("stopped"))
        VERIFY_STATE["error"] = None
    except Exception as e:
        VERIFY_STATE["error"] = f"{type(e).__name__}: {e}"
    finally:
        VERIFY_STATE["running"] = False
        VERIFY_STATE["finished_at"] = datetime.now().isoformat(timespec="seconds")


@app.post("/api/verify/run")
def verify_run():
    data = request.get_json(silent=True) or {}
    with _verify_lock:
        if VERIFY_STATE["running"]:
            return jsonify({"errors": ["A verification run is already in progress"]}), 409
        VERIFY_STATE.update(running=True, done=0, total=0, current=None,
                            verdicts={}, error=None, stopped=None, finished_at=None,
                            started_at=datetime.now().isoformat(timespec="seconds"))
    thread = threading.Thread(
        target=_run_verification,
        args=(data.get("limit"), data.get("category")),
        daemon=True)
    thread.start()
    return jsonify(VERIFY_STATE), 202


@app.get("/api/verify/status")
def verify_status():
    return jsonify(VERIFY_STATE)


@app.get("/api/review")
def review_queue():
    conn = get_conn()
    rows = conn.execute("""
        SELECT c.*, r.org_name, r.category, r.website AS current_website
        FROM verification_checks c JOIN resources r ON r.id = c.resource_id
        WHERE c.review_status = 'pending'
        ORDER BY CASE c.verdict
            WHEN 'possibly_inactive' THEN 0 WHEN 'changed' THEN 1
            WHEN 'unreachable' THEN 2 WHEN 'error' THEN 3
            WHEN 'needs_manual_check' THEN 4 ELSE 5 END, c.checked_at DESC
    """).fetchall()
    items = []
    for row in rows:
        item = dict(row)
        changes = json.loads(item.pop("proposed_changes") or "{}")
        resource = conn.execute("SELECT * FROM resources WHERE id = ?",
                                (item["resource_id"],)).fetchone()
        item["changes"] = [
            {"field": f, "old": resource[f], "new": v} for f, v in changes.items()
        ]
        items.append(item)
    conn.close()
    return jsonify(items)


def _close_check(check_id, action, overrides=None):
    """Shared approve/reject: returns (response, status).

    overrides: on approve, the exact {field: value} the staffer chose to apply
    (a subset of the AI's proposed changes, values possibly hand-edited). If
    None, all of the AI's proposed changes are applied (backward compatible).
    """
    conn = get_conn()
    check = conn.execute(
        "SELECT * FROM verification_checks WHERE id = ?", (check_id,)).fetchone()
    if check is None:
        conn.close()
        return jsonify({"errors": ["Check not found"]}), 404
    if check["review_status"] != "pending":
        conn.close()
        return jsonify({"errors": [f"Check already {check['review_status']}"]}), 409

    now = datetime.now().isoformat(timespec="seconds")
    today = date.today().isoformat()
    applied = []
    if action == "approve":
        proposed = json.loads(check["proposed_changes"] or "{}")
        raw = proposed if overrides is None else overrides
        sets, params = [], []
        for field, value in raw.items():
            if field not in verify_agent.PROPOSABLE_FIELDS:
                continue
            value = (str(value).strip() or None) if value is not None else None
            if field == "status" and value not in ("active", "needs_review", "inactive"):
                conn.close()
                return jsonify({"errors": ["status must be active, needs_review or inactive"]}), 400
            sets.append(f"{field} = ?")
            params.append(value)
            applied.append(field)
        # A staffer reviewing against the page counts as a verification, even
        # if they applied none of the AI's changes (current values confirmed).
        if check["verdict"] in ("confirmed", "changed", "possibly_inactive"):
            note = f"Staff reviewed AI check ({check['verdict']}, {today})"
            if applied:
                note += "; applied " + ", ".join(applied)
            sets += ["last_verified_date = ?", "verification_note = ?"]
            params += [today, note]
        if sets:
            sets.append("updated_at = ?")
            params.append(now)
            conn.execute(f"UPDATE resources SET {', '.join(sets)} WHERE id = ?",
                         params + [check["resource_id"]])
        new_status = "approved"
    else:
        new_status = "rejected"

    conn.execute(
        "UPDATE verification_checks SET review_status = ?, reviewed_at = ? WHERE id = ?",
        (new_status, now, check_id))
    conn.commit()
    conn.close()
    return jsonify({"ok": True, "review_status": new_status, "applied": applied}), 200


@app.post("/api/review/<int:check_id>/approve")
def approve_check(check_id):
    data = request.get_json(silent=True) or {}
    # 'changes' present (even {}) means "apply exactly these"; absent means "apply all".
    overrides = data["changes"] if isinstance(data.get("changes"), dict) else None
    return _close_check(check_id, "approve", overrides=overrides)


@app.post("/api/review/<int:check_id>/reject")
def reject_check(check_id):
    return _close_check(check_id, "reject")


# ---------- Phase 3: matching ----------

@app.post("/api/match")
def match():
    """Anonymous matching. The profile is used for this response only -
    nothing from the request body is stored or logged (PRD privacy rule)."""
    data = request.get_json(force=True, silent=True) or {}
    profile = {
        "needs": [c for c in (data.get("needs") or []) if isinstance(c, str)][:20],
        "area": str(data.get("area") or "").strip()[:80],
        "needs_spanish_or_bilingual": bool(data.get("needs_spanish")),
        "circumstances": [str(c)[:40] for c in (data.get("circumstances") or [])][:12],
        "notes": str(data.get("notes") or "").strip()[:1000],
    }
    use_ai = (data.get("mode") or "ai") != "basic"
    return jsonify(matcher.find_matches(profile, use_ai=use_ai))


@app.post("/api/share")
def create_share():
    """Save a match list under a unique link. Stores resource info + the
    AI's per-resource reason only - never anything about the family."""
    data = request.get_json(force=True, silent=True) or {}
    items = data.get("matches") or []
    conn = get_conn()
    payload = []
    for it in items[:20]:
        if not isinstance(it, dict):
            continue
        row = conn.execute("SELECT * FROM resources WHERE id = ?",
                           (it.get("id"),)).fetchone()
        if row:  # rebuild from the DB so the link always carries clean data
            payload.append(matcher._decorate(dict(row),
                                             str(it.get("reason") or "")[:300]))
    if not payload:
        conn.close()
        return jsonify({"errors": ["Nothing to share - run a match first"]}), 400
    token = secrets.token_urlsafe(8)
    conn.execute("INSERT INTO shared_lists (token, payload, created_at) VALUES (?, ?, ?)",
                 (token, json.dumps(payload, ensure_ascii=False),
                  datetime.now().isoformat(timespec="seconds")))
    conn.commit()
    conn.close()
    return jsonify({"url": request.host_url.rstrip("/") + "/l/" + token,
                    "token": token}), 201


@app.get("/l/<token>")
def shared_list(token):
    conn = get_conn()
    row = conn.execute("SELECT * FROM shared_lists WHERE token = ?",
                       (token,)).fetchone()
    conn.close()
    if row is None:
        return render_template("share.html", groups=None, created=None), 404
    matches = json.loads(row["payload"])
    groups, order = {}, []
    for m in matches:
        if m["category"] not in groups:
            groups[m["category"]] = []
            order.append(m["category"])
    for m in matches:
        groups[m["category"]].append(m)
    return render_template("share.html",
                           groups=[(c, groups[c]) for c in order],
                           created=row["created_at"][:10])


@app.get("/export.xlsx")
def export_xlsx():
    conn = get_conn()
    rows = conn.execute(
        "SELECT * FROM resources ORDER BY category, org_name COLLATE NOCASE").fetchall()
    conn.close()
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    by_category = {}
    for row in rows:
        by_category.setdefault(row["category"], []).append(row)
    for category, items in by_category.items():
        ws = wb.create_sheet(title=category[:31])
        ws.append(list(EXPORT_COLUMNS))
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for item in items:
            ws.append([item[c] for c in EXPORT_COLUMNS])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    name = f"community_resources_{datetime.now():%Y-%m-%d}.xlsx"
    return send_file(buf, as_attachment=True, download_name=name,
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


if __name__ == "__main__":
    # Local dev only. In production the app is served by gunicorn (see Procfile),
    # which never enables debug. Set FLASK_DEBUG=1 to turn on the reloader locally.
    import os
    app.run(debug=os.environ.get("FLASK_DEBUG") == "1",
            port=int(os.environ.get("PORT", 5057)))
